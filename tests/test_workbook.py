import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
from cmf.workbook import write_workbook,recalculate,set_column_width


class WorkbookTests(unittest.TestCase):
    def test_column_width_splits_large_group_without_expansion(self):
        sheet=openpyxl.Workbook().active
        dim=sheet.column_dimensions['M']
        dim.min,dim.max,dim.width=13,16384,10
        set_column_width(sheet,13,19)
        self.assertEqual(sheet.column_dimensions['M'].width,19)
        self.assertEqual(sheet.column_dimensions['M'].max,13)
        self.assertEqual(sheet.column_dimensions['N'].max,16384)
        self.assertEqual(sheet.column_dimensions['N'].width,10)
        self.assertEqual(len(sheet.column_dimensions),2)

    def test_conflict_and_missing_input_guards(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            wb=openpyxl.Workbook()
            wb.active.title='TAF_G1'
            wb.create_sheet('TAF_G3')
            wb['TAF_G1']['A51']='STAR'
            wb['TAF_G3']['B1']='STAR'
            wb['TAF_G1']['C54']=123
            wb['TAF_G3']['Q9']='=SUM(Q4:Q8)'
            wb['TAF_G3']['Q11']='=Q10'
            wb['TAF_G3']['Q12']='=Q9+Q11'
            source=root/'source.xlsx'
            wb.save(source)
            config=json.loads(Path('config/star.json').read_text())
            config['extend_template']=False
            r=dict(value=456,approved=True)
            annual={2025:dict(records={'AC12':r},validation={'eligible':['AC12'],'dependencies':{}},branches={})}
            with patch('cmf.workbook.recalculate'):
                writes,issues=write_workbook(source,root/'result.xlsx',config,annual)
            actual=openpyxl.load_workbook(root/'result.xlsx')
            self.assertEqual(actual['TAF_G1']['C54'].value,123)
            self.assertTrue(any(i['type']=='existing_cell_conflict' for i in issues))
            self.assertIn('COUNT',actual['TAF_G3']['Q9'].value)
            self.assertEqual(openpyxl.load_workbook(source)['TAF_G3']['Q9'].value,'=SUM(Q4:Q8)')

    @unittest.skipUnless(__import__('shutil').which('libreoffice'),'LibreOffice integration requires Docker')
    def test_real_recalculation_missing_zero_and_preservation(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'recalc.xlsx'
            wb=openpyxl.Workbook()
            s=wb.active
            s['A1']=10
            s['A2']=0
            s['B1']='=IF(COUNT(A1,A2)=2,SUM(A1,A2),"")'
            s['B2']='=IF(COUNT(A1,A3)=2,SUM(A1,A3),"")'
            wb.save(path)
            recalculate(path)
            values=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(values.active['B1'].value,10)
            self.assertIn(values.active['B2'].value,(None,''))
            self.assertEqual(openpyxl.load_workbook(path).active['B1'].value,s['B1'].value)
