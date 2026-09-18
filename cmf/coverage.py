"""Measure target-cell coverage, not number of writes or project progress."""
from collections import Counter
import openpyxl
from openpyxl.utils import get_column_letter


def measure(path,config,years,writes,issues):
    workbook=openpyxl.load_workbook(path,data_only=True)
    written={(w['sheet'],w['cell']) for w in writes}
    conflicts={(i['sheet'],i['cell']) for i in issues if i['type']=='existing_cell_conflict'}
    cells=[]
    for year in years:
        targets=[]
        for mapping,offset in [('asset_rows',2028),('liability_rows',2036)]:
            targets.extend(('TAF_G1',f'{get_column_letter(offset-year)}{row}',code)
                           for code,row in config[mapping].items())
        col=get_column_letter(year-2008)
        targets.extend(('TAF_G3',f'{col}{row}',code) for code,row in config['branch_rows'].items())
        targets.extend(('TAF_G3',f'{col}{row}',f'TOTAL_{row}') for row in [9,11,12])
        for sheet,cell,code in targets:
            value=workbook[sheet][cell].value
            status=('conflict' if (sheet,cell) in conflicts else
                    'missing_or_blocked' if value in (None,'') else
                    'populated_this_run' if (sheet,cell) in written else 'existing_not_revalidated')
            cells.append(dict(year=year,sheet=sheet,cell=cell,code=code,status=status))
    summary=[]
    for year in years:
        for sheet in ['TAF_G1','TAF_G3']:
            selected=[c for c in cells if c['year']==year and c['sheet']==sheet]
            counts=dict(Counter(c['status'] for c in selected))
            summary.append(dict(year=year,sheet=sheet,expected_cells=len(selected),counts=counts))
    return dict(basis='Configured target cells including totals; no assumption that unpublished rows are zero. STAR_Details excluded.',
                summary=summary,cells=cells)
