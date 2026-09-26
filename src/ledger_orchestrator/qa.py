"""Developer-only visual inspection of an output workbook."""
import argparse
import subprocess
import tempfile
from pathlib import Path
import openpyxl
import pypdfium2


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('workbook',type=Path)
    args=parser.parse_args()
    with tempfile.TemporaryDirectory() as temp:
        tmp=Path(temp)
        for sheet,area in [('TAF_G1','A51:O99'),('TAF_G3','A1:Q12')]:
            wb=openpyxl.load_workbook(args.workbook)
            for name in list(wb.sheetnames):
                # Keep dependent sheets for recalculation; deleting STAR_Details
                # makes otherwise valid cross-sheet totals disappear in previews.
                wb[name].sheet_state='visible' if name==sheet else 'hidden'
            wb.active=wb.sheetnames.index(sheet)
            s=wb[sheet]
            s.print_area=area
            s.page_setup.orientation='landscape'
            s.page_setup.paperSize=s.PAPERSIZE_A3
            s.sheet_properties.pageSetUpPr.fitToPage=True
            s.page_setup.fitToWidth=1
            s.page_setup.fitToHeight=1
            target=tmp/f'{sheet}.xlsx'
            wb.save(target)
            subprocess.run(['libreoffice',f'-env:UserInstallation={(tmp/"profile").as_uri()}',
                            '--headless','--convert-to','pdf','--outdir',str(tmp),str(target)],check=True,capture_output=True,timeout=120)
            pdf=pypdfium2.PdfDocument(tmp/f'{sheet}.pdf')
            try:
                bitmap=pdf[0].render(scale=1.6)
                image=bitmap.to_pil()
                image.save(args.workbook.parent/f'preview_{sheet}.png')
                image.close()
                bitmap.close()
            finally:
                pdf.close()


if __name__=='__main__':
    main()
